import fitz
from docx import Document
import openpyxl
from pathlib import Path
import tempfile
import os

def extract_pdf(path):
    doc = fitz.open(path)
    return "\n".join(page.get_text() for page in doc)

def extract_docx(path):
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def extract_xlsx(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    text = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join(str(c) for c in row if c is not None)
            if row_text.strip():
                text.append(row_text)
    return "\n".join(text)

def extract_text_from_file(uploaded_file):
    ext = Path(uploaded_file.name).suffix.lower()
    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        tmp.write(uploaded_file.read())
        tmp_path = tmp.name
    try:
        if ext == ".pdf":
            return extract_pdf(tmp_path)
        elif ext == ".docx":
            return extract_docx(tmp_path)
        elif ext == ".xlsx":
            return extract_xlsx(tmp_path)
        else:
            return None
    finally:
        os.unlink(tmp_path)